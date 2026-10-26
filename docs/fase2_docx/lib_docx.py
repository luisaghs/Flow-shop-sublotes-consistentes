# -*- coding: utf-8 -*-
"""
Utilidades para escribir la Fase 2 dentro del borrador del trabajo de grado,
reutilizando sus estilos (Normal, heading 2, Titulo 3, caption, List Paragraph,
Table Grid) y con ecuaciones nativas de Word (OMML) generadas con pandoc.

Marcado en línea de los textos:
  $...$   ecuación en línea (LaTeX)
  **...** negrita
  *...*   cursiva
  [[...]] texto resaltado en amarillo (pendiente de confirmar por las autoras/es)
"""
import copy
import os
import re
import subprocess
import tempfile
import zipfile

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from lxml import etree

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
ANCHO_TEXTO = 8838            # ancho útil de la página [dxa] (21,59 cm − 2 × 3 cm)
PATRON = re.compile(r'(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|\[\[[^\]]+\]\])')


# ---------------------------------------------------------------------------
# Ecuaciones: LaTeX → OMML con pandoc (una sola llamada para todas)
# ---------------------------------------------------------------------------

def latex_a_omml(snippets):
    """Retorna {latex: elemento m:oMath} para todos los fragmentos."""
    unicos = list(dict.fromkeys(snippets))
    if not unicos:
        return {}
    with tempfile.TemporaryDirectory() as tmp:
        md = os.path.join(tmp, 'ec.md')
        with open(md, 'w', encoding='utf-8') as fh:
            fh.write('\n\n'.join(f'${s}$' for s in unicos) + '\n')
        salida = os.path.join(tmp, 'ec.docx')
        subprocess.run(['pandoc', md, '-f', 'markdown+tex_math_dollars', '-o', salida], check=True)
        with zipfile.ZipFile(salida) as z:
            raiz = etree.fromstring(z.read('word/document.xml'))
    parrafos = raiz.findall(f'.//{{{W_NS}}}body/{{{W_NS}}}p')
    res = {}
    for s, p in zip(unicos, parrafos):
        om = p.find(f'.//{{{M_NS}}}oMath')
        if om is None:
            raise ValueError(f'pandoc no convirtió la ecuación: {s}')
        res[s] = om
    if len(res) != len(unicos):
        raise ValueError('No coinciden las ecuaciones convertidas.')
    return res


class Escritor:
    """Agrega contenido al final del documento con los estilos del borrador."""

    def __init__(self, documento, omml, primer_figura=None):
        self.d = documento
        self.omml = omml
        self.tabla_n = 0
        self.figura_n = 0
        self.primer_figura = primer_figura        # número forzado de la primera figura
        self._numeracion = self.d.part.numbering_part.element
        self._siguiente_num = max(int(n.get(qn('w:numId')))
                                  for n in self._numeracion.findall(qn('w:num'))) + 1

    # ── texto en línea ──
    def _runs(self, p, texto, tam=None, negrita=False):
        for trozo in PATRON.split(texto):
            if not trozo:
                continue
            if trozo.startswith('$') and trozo.endswith('$') and len(trozo) > 2:
                p._p.append(copy.deepcopy(self.omml[trozo[1:-1]]))
                continue
            r = None
            if trozo.startswith('**') and trozo.endswith('**'):
                r = p.add_run(trozo[2:-2]); r.bold = True
            elif trozo.startswith('[[') and trozo.endswith(']]'):
                r = p.add_run(trozo[2:-2]); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif trozo.startswith('*') and trozo.endswith('*') and len(trozo) > 2:
                r = p.add_run(trozo[1:-1]); r.italic = True
            else:
                r = p.add_run(trozo)
            if negrita:
                r.bold = True
            if tam:
                r.font.size = Pt(tam)

    # ── bloques ──
    def h2(self, texto):
        self.d.add_paragraph(texto, style='Heading 2')

    def h3(self, texto):
        self.d.add_paragraph(texto, style='Titulo 3')

    def h4(self, texto):
        p = self.d.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._runs(p, texto, negrita=True)

    def p(self, texto):
        p = self.d.add_paragraph()
        self._runs(p, texto)
        return p

    def _lista(self, items, num_id):
        for t in items:
            p = self.d.add_paragraph(style='List Paragraph')
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement('w:numPr')
            ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), '0')
            nid = OxmlElement('w:numId'); nid.set(qn('w:val'), str(num_id))
            numPr.append(ilvl); numPr.append(nid); pPr.append(numPr)
            self._runs(p, t)

    def viñetas(self, items):
        self._lista(items, 3)               # lista con viñetas del borrador

    def numerada(self, items):
        """Lista numerada que reinicia en 1 (nueva instancia de la lista decimal)."""
        nid = self._siguiente_num
        self._siguiente_num += 1
        num = etree.SubElement(self._numeracion, qn('w:num'))
        num.set(qn('w:numId'), str(nid))
        a = etree.SubElement(num, qn('w:abstractNumId')); a.set(qn('w:val'), '10')
        ov = etree.SubElement(num, qn('w:lvlOverride')); ov.set(qn('w:ilvl'), '0')
        st = etree.SubElement(ov, qn('w:startOverride')); st.set(qn('w:val'), '1')
        self._lista(items, nid)

    # ── rótulos ──
    def _rotulo(self, tipo, titulo, reiniciar=None):
        p = self.d.add_paragraph(style='Caption')
        p.paragraph_format.keep_with_next = True

        def run(texto=None):
            r = p.add_run(texto) if texto is not None else p.add_run()
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0, 0, 0)
            return r
        run(f'{tipo} ')
        if tipo == 'Tabla':
            self.tabla_n += 1; n = self.tabla_n
        else:
            self.figura_n += 1; n = self.figura_n
        instr = f' SEQ {tipo} \\* ARABIC ' + (f'\\r {reiniciar} ' if reiniciar else '')
        for tipo_f, texto in (('begin', None), ('instr', instr), ('separate', None),
                              ('num', str(n)), ('end', None)):
            r = run()
            if tipo_f == 'instr':
                it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = texto
                r._r.append(it)
            elif tipo_f == 'num':
                r.text = texto
            else:
                fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo_f); r._r.append(fc)
        run(f'. {titulo}')
        return n

    # ── tablas ──
    @staticmethod
    def _bordes_celda(celda, arriba=False, abajo=False):
        tcPr = celda._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        for lado, activo in (('top', arriba), ('bottom', abajo)):
            if activo:
                e = OxmlElement(f'w:{lado}')
                for k, v in (('val', 'single'), ('sz', '4'), ('space', '0'), ('color', 'auto')):
                    e.set(qn(f'w:{k}'), v)
                b.append(e)
        tcPr.append(b)

    @staticmethod
    def _tabla_base(tabla, anchos, bordes_tabla=True):
        tbl = tabla._tbl
        tblPr = tbl.tblPr
        jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
        bordes = OxmlElement('w:tblBorders')
        lados = ['left', 'right', 'insideH', 'insideV'] + ([] if bordes_tabla else ['top', 'bottom'])
        for lado in lados:
            e = OxmlElement(f'w:{lado}')
            for k, v in (('val', 'none'), ('sz', '0'), ('space', '0'), ('color', 'auto')):
                e.set(qn(f'w:{k}'), v)
            bordes.append(e)
        tblPr.append(bordes)
        tabla.autofit = False
        for k, w in enumerate(anchos):
            for celda in tabla.columns[k].cells:
                celda.width = docx.shared.Twips(w)
        grid = tbl.tblGrid
        for k, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(anchos[k]))

    def tabla(self, titulo, encabezado, filas, anchos=None, alineacion=None, tam=11,
              negrita_filas=()):
        """Tabla con formato APA del borrador: líneas horizontales, texto de 11 pt."""
        self._rotulo('Tabla', titulo)
        ncol = len(encabezado)
        if anchos is None:
            anchos = [ANCHO_TEXTO // ncol] * ncol
        if alineacion is None:
            alineacion = ['l'] + ['c'] * (ncol - 1)
        t = self.d.add_table(rows=1 + len(filas), cols=ncol)
        t.style = self.d.styles['Table Grid']
        self._tabla_base(t, anchos)
        al = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'c': WD_ALIGN_PARAGRAPH.CENTER,
              'r': WD_ALIGN_PARAGRAPH.RIGHT}
        for k, texto in enumerate(encabezado):
            c = t.rows[0].cells[k]
            self._bordes_celda(c, arriba=True, abajo=True)
            par = c.paragraphs[0]; par.alignment = al[alineacion[k]]
            self._runs(par, texto, tam=tam, negrita=True)
        for i, fila in enumerate(filas, start=1):
            for k, texto in enumerate(fila):
                c = t.rows[i].cells[k]
                c.vertical_alignment = docx.enum.table.WD_ALIGN_VERTICAL.CENTER
                par = c.paragraphs[0]; par.alignment = al[alineacion[k]]
                self._runs(par, str(texto), tam=tam, negrita=(i - 1) in negrita_filas)
        # La fila de encabezado se repite si la tabla cambia de página
        trPr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement('w:tblHeader'); trPr.append(h)
        self.d.add_paragraph()
        return t

    def ecuacion(self, latex, dominio='', numero=None):
        """Ecuación numerada en una tabla sin bordes: ecuación | dominio | (n)."""
        t = self.d.add_table(rows=1, cols=3)
        t.style = self.d.styles['Table Grid']
        anchos = [5200, 2900, 738]
        self._tabla_base(t, anchos, bordes_tabla=False)
        c0, c1, c2 = t.rows[0].cells
        for c in (c0, c1, c2):
            c.vertical_alignment = docx.enum.table.WD_ALIGN_VERTICAL.CENTER
        if not dominio:                       # sin dominio: la ecuación usa ambas columnas
            c0 = c0.merge(c1)
        p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para = etree.SubElement(p0._p, f'{{{M_NS}}}oMathPara')
        ppr = etree.SubElement(para, f'{{{M_NS}}}oMathParaPr')
        jc = etree.SubElement(ppr, f'{{{M_NS}}}jc'); jc.set(f'{{{M_NS}}}val', 'center')
        para.append(copy.deepcopy(self.omml[latex]))
        # El dominio puede ir en varias líneas separadas por ";;"
        for k, parte in enumerate([x.strip() for x in dominio.split(';;')] if dominio else []):
            par = c1.paragraphs[0] if k == 0 else c1.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            self._runs(par, parte, tam=11)
        for par in (p0, c2.paragraphs[0]):
            par.paragraph_format.space_before = Pt(4)
            par.paragraph_format.space_after = Pt(4)
        p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if numero is not None:
            r = p2.add_run(f'({numero})')
        return t

    def figura(self, titulo, ruta, ancho_cm=15.0):
        reiniciar = self.primer_figura if (self.figura_n == 0 and self.primer_figura) else None
        self._rotulo('Figura', titulo, reiniciar)
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_together = True
        p.add_run().add_picture(ruta, width=Cm(ancho_cm))
        return self.figura_n

    def codigo(self, titulo, lineas):
        """Pseudocódigo en un recuadro con líneas superior e inferior."""
        reiniciar = self.primer_figura if (self.figura_n == 0 and self.primer_figura) else None
        self._rotulo('Figura', titulo, reiniciar)
        t = self.d.add_table(rows=1, cols=1)
        t.style = self.d.styles['Table Grid']
        self._tabla_base(t, [ANCHO_TEXTO])
        c = t.rows[0].cells[0]
        primero = True
        for ln in lineas:
            par = c.paragraphs[0] if primero else c.add_paragraph()
            primero = False
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
            sangria = len(ln) - len(ln.lstrip(' '))
            par.paragraph_format.left_indent = Cm(0.45 * (sangria // 2))
            self._runs(par, ln.strip(), tam=10)
        self.d.add_paragraph()
        return self.figura_n


ORDEN_HIJOS = {
    # Orden exigido por el esquema OOXML
    f'{{{M_NS}}}rPr': ['lit', 'nor', 'scr', 'sty', 'brk', 'aln'],
    f'{{{M_NS}}}dPr': ['begChr', 'sepChr', 'endChr', 'grow', 'shp', 'ctrlPr'],
    f'{{{W_NS}}}tblPr': ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
                         'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd',
                         'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
                         'tblCaption', 'tblDescription'],
    f'{{{W_NS}}}tblBorders': ['top', 'left', 'start', 'bottom', 'right', 'end', 'insideH', 'insideV'],
    f'{{{W_NS}}}tcBorders': ['top', 'left', 'start', 'bottom', 'right', 'end', 'insideH', 'insideV',
                             'tl2br', 'tr2bl'],
    f'{{{W_NS}}}tcPr': ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd',
                        'noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark'],
}


def normalizar_orden(documento):
    """Reordena los hijos de los elementos con orden fijo según el esquema."""
    raiz = documento.element.body
    for etiqueta, orden in ORDEN_HIJOS.items():
        pos = {nombre: k for k, nombre in enumerate(orden)}
        for el in raiz.iter(etiqueta):
            hijos = list(el)
            ordenados = sorted(hijos, key=lambda h: pos.get(etree.QName(h).localname, len(orden)))
            if ordenados != hijos:
                for h in hijos:
                    el.remove(h)
                for h in ordenados:
                    el.append(h)


def resolver_referencias(bloques, tabla_ini=1, figura_ini=1):
    """Numera tablas y figuras con etiqueta y reemplaza {T:etiqueta} y {F:etiqueta}."""
    nums, t, f = {}, tabla_ini, figura_ini
    for b in bloques:
        if b[0] == 'tabla':
            kw = b[4] if len(b) > 4 else {}
            if kw.get('etiqueta'):
                nums['T:' + kw['etiqueta']] = t
            t += 1
        elif b[0] == 'figura':
            if len(b) > 4:
                nums['F:' + b[4]] = f
            f += 1
        elif b[0] == 'codigo':
            if len(b) > 3:
                nums['F:' + b[3]] = f
            f += 1

    def sub(x):
        if isinstance(x, str):
            def rep(m):
                if m.group(1) not in nums:
                    raise KeyError(f'Referencia sin destino: {m.group(1)}')
                return str(nums[m.group(1)])
            return re.sub(r'\{([TF]:[a-z0-9_]+)\}', rep, x)
        if isinstance(x, list):
            return [sub(y) for y in x]
        if isinstance(x, tuple):
            return tuple(sub(y) for y in x)
        if isinstance(x, dict):
            return {k: sub(v) for k, v in x.items()}
        return x
    return [sub(b) for b in bloques]


SIMBOLOS = [
    (re.compile(r'ΣT_j'), lambda m: '$\\sum T_j$'),
    (re.compile(r'ΣE_j'), lambda m: '$\\sum E_j$'),
    (re.compile(r'C_max'), lambda m: '$C_{\\max}$'),
    (re.compile(r'T_max'), lambda m: '$T_{\\max}$'),
    (re.compile(r'(?<![\w$])p_([cm])(?![\w])'), lambda m: f'$p_{m.group(1)}$'),
    (re.compile(r'(?<![\w$\\])([dQrCTE])_j(?![\w])'), lambda m: f'${m.group(1)}_j$'),
]


def _matematizar_texto(t):
    if not isinstance(t, str):
        return t
    for patron, reemplazo in SIMBOLOS:          # cada patrón respeta las ecuaciones previas
        partes = re.split(r'(\$[^$]+\$)', t)
        t = ''.join(p if p.startswith('$') else patron.sub(reemplazo, p) for p in partes)
    return t


def matematizar(bloques):
    """Convierte símbolos escritos como texto (ΣT_j, C_max, d_j…) en ecuaciones en línea.
    No toca el pseudocódigo ni los rótulos."""
    res = []
    for b in bloques:
        tipo = b[0]
        if tipo in ('p', 'h4'):
            b = (tipo, _matematizar_texto(b[1])) + tuple(b[2:])
        elif tipo in ('viñetas', 'numerada'):
            b = (tipo, [_matematizar_texto(x) for x in b[1]])
        elif tipo == 'tabla':
            b = (tipo, b[1], [_matematizar_texto(x) for x in b[2]],
                 [[_matematizar_texto(x) for x in fila] for fila in b[3]]) + tuple(b[4:])
        res.append(b)
    return res


def recolectar_latex(bloques):
    """Busca todos los fragmentos $...$ y las ecuaciones de bloque."""
    s = []

    def de_texto(t):
        for tr in PATRON.split(str(t)):
            if tr.startswith('$') and tr.endswith('$') and len(tr) > 2:
                s.append(tr[1:-1])
    for b in bloques:
        tipo = b[0]
        if tipo in ('p', 'h4'):
            de_texto(b[1])
        elif tipo in ('viñetas', 'numerada'):
            for t in b[1]:
                de_texto(t)
        elif tipo == 'tabla':
            for t in b[2]:
                de_texto(t)
            for fila in b[3]:
                for t in fila:
                    de_texto(t)
        elif tipo == 'ecuacion':
            s.append(b[1]); de_texto(b[2])
        elif tipo == 'codigo':
            for t in b[2]:
                de_texto(t)
    return s


def escribir(escritor, bloques):
    for b in bloques:
        tipo, args = b[0], b[1:]
        if tipo == 'h2':
            escritor.h2(*args)
        elif tipo == 'h3':
            escritor.h3(*args)
        elif tipo == 'h4':
            escritor.h4(*args)
        elif tipo == 'p':
            escritor.p(*args)
        elif tipo == 'viñetas':
            escritor.viñetas(*args)
        elif tipo == 'numerada':
            escritor.numerada(*args)
        elif tipo == 'tabla':
            titulo, enc, filas = args[:3]
            kw = dict(args[3]) if len(args) > 3 else {}
            kw.pop('etiqueta', None)
            escritor.tabla(titulo, enc, filas, **kw)
        elif tipo == 'ecuacion':
            escritor.ecuacion(*args)
        elif tipo == 'figura':
            escritor.figura(*args[:3])
        elif tipo == 'codigo':
            escritor.codigo(*args[:2])
        elif tipo == 'vacio':
            escritor.d.add_paragraph()
        else:
            raise ValueError(tipo)


def quitar_desde(documento, texto_encabezado):
    """Elimina el encabezado indicado y todo lo que le sigue (excepto sectPr)."""
    cuerpo = documento.element.body
    inicio = None
    for el in cuerpo.iterchildren():
        if el.tag == qn('w:p') and \
                docx.text.paragraph.Paragraph(el, documento).text.strip() == texto_encabezado:
            inicio = el
            break
    if inicio is None:
        raise ValueError(f'No se encontró el encabezado "{texto_encabezado}".')
    borrar, activo = [], False
    for el in list(cuerpo.iterchildren()):
        if el is inicio:
            activo = True
        if activo and el.tag != qn('w:sectPr'):
            borrar.append(el)
    for el in borrar:
        cuerpo.remove(el)
    return len(borrar)
